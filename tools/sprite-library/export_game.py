"""Export the Sōzō sprite library as game assets.

Usage (from the repo root):
    python tools/sprite-library/export_game.py

Reads documents/sozo-sprite-library.xlsx as the manifest (it lists every included entry with its
Sōzō ID, stats and the reference ID it covers), re-draws each entry with this generator, and writes:

  src/generated/library.ts      typed data for every sheet (blocks, items, creatures, effects...)
  src/generated/libraryArt.ts   the SVG drawing for every 2D sprite, keyed "<kind>:<id>"
  src/generated/blockIds.json   stable numeric ids for library blocks (append-only)
  public/library/blocks.png     flat face textures for blocks (TILE_PX squares)
  public/library/models/*.png   the 3D-rendered sprites (creatures, vehicles, structures)

Block faces are flat textures (what chunks need), not the isometric icons: each block's main box is
re-painted face by face from the same materials and brush patterns. Blocks that are not boxes
(plants, torches, lanterns...) get their whole sprite as one "sprite" tile instead.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import cairosvg  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402

import classify_blocks as cb  # noqa: E402
import creatures as creature_items  # noqa: E402
import geometry  # noqa: E402
import items_food as food_items  # noqa: E402
import items_fx as fx  # noqa: E402
import items_gear as gear_items  # noqa: E402
import items_materials as mat_items  # noqa: E402
import plants  # noqa: E402
import shapes  # noqa: E402
from geometry import Rand  # noqa: E402

ROOT = HERE.parent.parent
XLSX = ROOT / "documents" / "sozo-sprite-library.xlsx"
GEN = ROOT / "src" / "generated"
PUB = ROOT / "public" / "library"
TILE_PX = 32
ATLAS_COLS = 32

# The game's hand-written blocks keep their ids (src/world/blocks.ts); library blocks with the same
# key share that id and become the block's "Library art" look.
CORE_BLOCK_IDS = {
    "air": 0, "grass": 1, "dirt": 2, "stone": 3, "sand": 4, "water": 5, "log": 6, "leaves": 7,
    "planks": 8, "glass": 9, "cobblestone": 10, "coal_ore": 11, "iron_ore": 12, "brick": 13,
    "snow": 14, "crystal_ore": 15, "gravel": 16, "bedrock": 17, "wool": 18, "lamp": 19,
    "photo_block": 20,
    **{f"water_flow_{s}": 20 + s for s in range(1, 8)},
}
# Library block → core block it is the library look of (where the Sōzō id differs).
CORE_ALIAS = {
    "grass_block": "grass", "oak_log": "log", "oak_leaves": "leaves", "oak_planks": "planks",
    "bricks": "brick", "snow_block": "snow", "white_wool": "wool",
}


# --- reading the manifest -----------------------------------------------------------------------

def sheet_rows(wb, title: str) -> list[dict]:
    ws = wb[title]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:] if r[1]]


def camel(h: str) -> str:
    words = "".join(ch if ch.isalnum() else " " for ch in h.replace("Sōzō", "")).split()
    return words[0].lower() + "".join(w.title() for w in words[1:]) if words else h


def clean(v):
    if v in ("", "—", None):
        return None
    if v in ("Yes", "No"):
        return v == "Yes"
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def record(row: dict, skip=("Sprite",)) -> dict:
    out = {}
    for k, v in row.items():
        if k in skip:
            continue
        if k == "Sōzō ID":
            out["id"] = v
        elif k == "Replaces (reference ID)":
            out["ref"] = v
        else:
            out[camel(k)] = clean(v)
    return out


# --- block faces --------------------------------------------------------------------------------

_captured: list = []
_real_boxes_svg = geometry.boxes_svg


def _capturing_boxes_svg(boxes, rand):
    _captured.append(list(boxes))
    return _real_boxes_svg(boxes, rand)


for _mod in (geometry, shapes, plants, cb):
    if hasattr(_mod, "boxes_svg"):
        _mod.boxes_svg = _capturing_boxes_svg


def draw_with_boxes(draw):
    _captured.clear()
    svg = draw()
    boxes = [b for group in _captured for b in group]
    return svg, boxes


def face_svg(fill: str, painter, seed: str, opacity: float) -> str:
    op = f' fill-opacity="{opacity}"' if opacity < 1 else ""
    body = f'<rect x="0" y="0" width="16" height="16" fill="{fill}"{op}/>'
    if painter is not None:
        body += painter(16, 16, fill, Rand(seed)) or ""
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 16 16" width="16" height="16">{body}</svg>'


def sprite_svg(svg_body: str) -> str:
    """A whole sprite as a tile (plants, torches...): the icon box trimmed to its 0–100 content."""
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="-3 -3 106 106" width="106" height="106">{svg_body}</svg>'


class Atlas:
    def __init__(self):
        self.tiles: list[bytes] = []
        self.by_hash: dict[str, int] = {}

    def add(self, svg: str) -> int:
        png = cairosvg.svg2png(bytestring=svg.encode(), output_width=TILE_PX, output_height=TILE_PX)
        key = hashlib.sha1(Image.open(io.BytesIO(png)).convert("RGBA").tobytes()).hexdigest()
        if key not in self.by_hash:
            self.by_hash[key] = len(self.tiles)
            self.tiles.append(png)
        return self.by_hash[key]

    def save(self, path: Path) -> tuple[int, int]:
        rows = max(1, math.ceil(len(self.tiles) / ATLAS_COLS))
        img = Image.new("RGBA", (ATLAS_COLS * TILE_PX, rows * TILE_PX), (0, 0, 0, 0))
        for i, png in enumerate(self.tiles):
            t = Image.open(io.BytesIO(png)).convert("RGBA")
            img.paste(t, ((i % ATLAS_COLS) * TILE_PX, (i // ATLAS_COLS) * TILE_PX))
        img.save(path, optimize=True)
        return ATLAS_COLS, rows


def block_tiles(entry, atlas: Atlas) -> dict:
    svg, boxes = draw_with_boxes(entry.draw)
    if entry.shape not in ("plant", "custom") and boxes:
        # The biggest box carries the block's material (stairs, slabs and fences share it).
        bx = max(boxes, key=lambda b: (b.a1 - b.a0) * (b.b1 - b.b0) * (b.c1 - b.c0))
        top_c = bx.top_color or bx.color
        side_c = bx.side_color or bx.color
        seed = entry.src_id
        top = atlas.add(face_svg(top_c, bx.top, seed + ":top", bx.opacity))
        side = atlas.add(face_svg(side_c, bx.side, seed + ":side", bx.opacity))
        return {"top": top, "side": side, "bottom": top}, svg
    return {"sprite": atlas.add(sprite_svg(svg))}, svg


# --- export ---------------------------------------------------------------------------------------

def main() -> None:
    wb = load_workbook(XLSX, read_only=True)
    GEN.mkdir(parents=True, exist_ok=True)
    (PUB / "models").mkdir(parents=True, exist_ok=True)
    art: dict[str, str] = {}
    atlas = Atlas()
    problems: list[str] = []

    def put_art(key: str, svg: str) -> None:
        if key in art:
            problems.append(f"duplicate art key {key}")
        art[key] = svg

    # Blocks.
    blocks = []
    for row in sheet_rows(wb, "Blocks"):
        rec = record(row)
        e = cb.classify(rec["ref"], "")
        if e is None:
            problems.append(f"block {rec['id']}: classifier no longer covers {rec['ref']}")
            continue
        tiles, svg = block_tiles(e, atlas)
        rec["tiles"] = tiles
        # Same key as a core block only counts when it is the same kind of thing (the library's
        # "snow" is a thin layer, not the core snow block).
        same = rec["id"] in CORE_BLOCK_IDS and rec["shape"] == "cube" and rec["id"] not in CORE_ALIAS.values()
        rec["core"] = CORE_ALIAS.get(rec["id"], rec["id"] if same else None)
        put_art(f"block:{rec['id']}", svg)
        blocks.append(rec)

    # Materials, gear, food: drawings come from the batch catalogues, matched by reference id.
    draws = {src: d for src, (_, _, d) in mat_items.MATERIALS.items()}
    draws.update({g["src"]: g["draw"] for g in gear_items.gear_catalogue()})
    draws.update({f["src"]: f["draw"] for f in food_items.food_catalogue()})

    def item_sheet(title: str) -> list[dict]:
        out = []
        for row in sheet_rows(wb, title):
            rec = record(row)
            draw = draws.get(rec["ref"])
            if draw is None:
                problems.append(f"{title} {rec['id']}: no drawing for {rec['ref']}")
                continue
            put_art(f"item:{rec['id']}", draw())
            out.append(rec)
        return out

    materials = item_sheet("Materials")
    gear = item_sheet("Tools & Gear")
    food = item_sheet("Food")

    # Potions: the drawing depends on the brew colour and the form.
    potions = []
    for row in sheet_rows(wb, "Potions"):
        rec = record(row)
        ref = rec["ref"]
        form = {"Drink": "drink", "Splash": "splash", "Mist": "mist"}[rec["form"]]
        if "water" in ref:
            svg = fx.potion("#a9cfe0", "drink", 0.5)
        elif "awkward" in ref:
            svg = fx.potion("#9aa8c0", "drink")
        else:
            brew = ref.split('minecraft:')[1].split('"')[0]
            svg = fx.potion(fx.POTIONS[brew][2], form)
        put_art(f"item:{rec['id']}", svg)
        potions.append(rec)

    # Effects and particles.
    effects = []
    for row in sheet_rows(wb, "Effects"):
        rec = record(row)
        _, col, _, _, symbol = fx.EFFECTS[rec["ref"]]
        put_art(f"effect:{rec['id']}", fx.badge(col, symbol))
        effects.append(rec)
    particles = []
    for row in sheet_rows(wb, "Particles"):
        rec = record(row)
        put_art(f"particle:{rec['id']}", fx.PARTICLES[rec["ref"]][3]())
        particles.append(rec)

    # Creatures and structures are 3D models: copy their rendered sprites for icons.
    models_dir = HERE / "sprites3d"
    crit = {c["src"]: c for c in creature_items.creature_catalogue()}

    def model_sheet(title: str) -> list[dict]:
        out = []
        for row in sheet_rows(wb, title):
            rec = record(row)
            src = models_dir / f"{rec['ref']}.png"
            if src.exists():
                shutil.copyfile(src, PUB / "models" / f"{rec['id']}.png")
                rec["sprite"] = f"library/models/{rec['id']}.png"
            else:
                # A few objects (e.g. the item frame) are still 2D drawings.
                if rec["ref"] in crit:
                    put_art(f"creature:{rec['id']}", crit[rec["ref"]]["draw"]())
                else:
                    problems.append(f"{title} {rec['id']}: no sprite")
            out.append(rec)
        return out

    creatures = model_sheet("Creatures")
    structures = model_sheet("Structures")

    # Stable block ids: keep every id ever handed out; append new keys.
    ids_path = GEN / "blockIds.json"
    ids: dict[str, int] = json.loads(ids_path.read_text()) if ids_path.exists() else {}
    for k, v in CORE_BLOCK_IDS.items():
        if ids.get(k, v) != v:
            problems.append(f"block id for core key {k} changed")
        ids[k] = v
    next_id = max(ids.values()) + 1
    cores = [b["core"] for b in blocks if b["core"]]
    for k in {c for c in cores if cores.count(c) > 1}:
        problems.append(f"several library blocks claim core block {k}")
    for b in blocks:
        # The game block key: the core key it stands in for, else its own id (suffixed with its
        # shape if that id is already a core block's, e.g. the "snow" layer → "snow_layer").
        key = b["core"] or (b["id"] if b["id"] not in CORE_BLOCK_IDS else f"{b['id']}_{b['shape']}")
        b["key"] = key
        if key not in ids:
            ids[key] = next_id
            next_id += 1
        b["blockId"] = ids[key]
    ids_path.write_text(json.dumps(ids, indent=1) + "\n")

    cols, rows = atlas.save(PUB / "blocks.png")

    def ts_array(name: str, typ: str, data: list) -> str:
        body = ",\n".join("  " + json.dumps(d, ensure_ascii=False) for d in data)
        return f"export const {name}: readonly {typ}[] = [\n{body},\n];\n"

    header = (
        "// Generated by tools/sprite-library/export_game.py from documents/sozo-sprite-library.xlsx.\n"
        "// Do not edit by hand: change the generator and run `npm run sprites:export`.\n"
    )
    lib = header + (
        "import type {\n  LibBlock, LibMaterial, LibGear, LibFood, LibPotion, LibEffect, LibParticle, LibCreature, LibStructure,\n"
        "} from '../library/types';\n\n"
        f"export const LIB_TILE_PX = {TILE_PX};\nexport const LIB_ATLAS_COLS = {cols};\nexport const LIB_ATLAS_ROWS = {rows};\n"
        f"export const LIB_ATLAS_URL = 'library/blocks.png';\n\n"
    )
    lib += ts_array("LIB_BLOCKS", "LibBlock", blocks) + "\n"
    lib += ts_array("LIB_MATERIALS", "LibMaterial", materials) + "\n"
    lib += ts_array("LIB_GEAR", "LibGear", gear) + "\n"
    lib += ts_array("LIB_FOOD", "LibFood", food) + "\n"
    lib += ts_array("LIB_POTIONS", "LibPotion", potions) + "\n"
    lib += ts_array("LIB_EFFECTS", "LibEffect", effects) + "\n"
    lib += ts_array("LIB_PARTICLES", "LibParticle", particles) + "\n"
    lib += ts_array("LIB_CREATURES", "LibCreature", creatures) + "\n"
    lib += ts_array("LIB_STRUCTURES", "LibStructure", structures)
    (GEN / "library.ts").write_text(lib)

    art_ts = header + (
        "// Each value is the body of an SVG in the icon box (viewBox \"-6 -6 112 112\"); wrap it with\n"
        "// libSvg() from src/library/art.ts.\n\n"
        "export const LIBRARY_ART: Record<string, string> = {\n"
        + "".join(f"  {json.dumps(k)}: {json.dumps(v, ensure_ascii=False)},\n" for k, v in sorted(art.items()))
        + "};\n"
    )
    (GEN / "libraryArt.ts").write_text(art_ts)

    counts = dict(blocks=len(blocks), materials=len(materials), gear=len(gear), food=len(food), potions=len(potions),
                  effects=len(effects), particles=len(particles), creatures=len(creatures), structures=len(structures))
    print("exported", counts, "total", sum(counts.values()))
    print(f"atlas: {len(atlas.tiles)} unique tiles, {cols}x{rows}; art entries: {len(art)}; block ids up to {next_id - 1}")
    for p in problems:
        print("WARNING:", p)
    if problems:
        sys.exit(1)


if __name__ == "__main__":
    main()

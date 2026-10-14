"""Export the reference spreadsheet (names, IDs and facts only, no images) to JSON for build.py.

Usage: python export_reference.py <reference.xlsx> <out.json>
"""

import json
import sys

import openpyxl

SHEETS = ["Items", "Blocks", "Entities", "Effects", "Particles", "Structures"]


def main(src: str, out: str) -> None:
    wb = openpyxl.load_workbook(src, read_only=True)
    data = {}
    for sheet in SHEETS:
        ws = wb[sheet]
        top, sub = list(ws.iter_rows(min_row=2, max_row=3, values_only=True))
        headers = [(str(a) if a else "") + (" / " + str(b) if b else "") for a, b in zip(top, sub)]
        rows = []
        for r in ws.iter_rows(min_row=5, values_only=True):
            key = r[1] if sheet == "Particles" else r[2]
            if isinstance(key, str):
                rows.append({h: v for h, v in zip(headers, r) if h})
        data[sheet] = {"headers": headers, "rows": rows}
    json.dump(data, open(out, "w"), default=str)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
